#!/usr/bin/env python3
"""生成示例 Excel：samples/订单示例.xlsx。

覆盖典型数据：多 Sheet、中文表头、日期列、超长编号、空行、空 Sheet。
"""
import datetime
from pathlib import Path

from openpyxl import Workbook

ROOT = Path(__file__).resolve().parent.parent


def main():
    wb = Workbook()
    ws = wb.active
    ws.title = "订单表"
    ws.append(["订单号", "商品名称", "金额", "下单时间", "备注", "长编号"])
    goods = [
        ("手机壳", 29.9, "测试"),
        ("手机膜", 9.9, None),
        ("机械键盘", 199.0, "含手机支架"),
        ("显示器", 899.5, "新品"),
        ("鼠标垫", 19.9, "赠品"),
        ("手机充电器", 49.0, "热卖"),
        ("蓝牙耳机", 129.0, "手机周边"),
        ("显示器支架", 69.0, None),
    ]
    for i, (name, price, note) in enumerate(goods, start=1001):
        ws.append([i, name, price,
                   datetime.datetime(2026, 9, i - 1000, 10, 0, 0),
                   note, "1234567890123456%d" % (i % 10)])
    ws.append([None] * 6)  # 空行：导入时跳过
    ws.append([2001, "USB 数据线", 9.9, "2026-09-10", "手机配件", "12345678901234567"])

    ws2 = wb.create_sheet("客户表")
    ws2.append(["客户名称", "手机", "累计消费"])
    for name, phone, amount in [
            ("张三", "13800000000", 229.8),
            ("李四", "13911112222", 199.0),
            ("王五", None, 49.0)]:
        ws2.append([name, phone, amount])

    wb.create_sheet("草稿（空）")

    out = ROOT / "samples" / "订单示例.xlsx"
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print("已生成：%s" % out)


if __name__ == "__main__":
    main()
