"""测试数据构造：生成覆盖典型脏数据的示例 xlsx。"""
import datetime
from pathlib import Path

from openpyxl import Workbook


def make_sample_xlsx(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "订单表"
    ws.append(["订单号", "商品名称", "金额", "下单时间", "备注", "长编号"])
    ws.append([1001, "手机壳", 29.9,
               datetime.datetime(2026, 9, 1, 10, 0, 0), "测试", "12345678901234567"])
    ws.append([1002, "手机膜", 9.9,
               datetime.datetime(2026, 9, 2), None, "12345678901234568"])
    ws.append([1003, "机械键盘", 199.0, "2026-09-03", "含手机支架", "12345678901234569"])
    ws.append([None] * 6)  # 空行
    ws.append([1004, "显示器", 899.5, "2026/09/05 08:30:00", "新品", "12345678901234570"])

    ws2 = wb.create_sheet("客户表")
    ws2.append(["客户名称", "手机", "消费"])
    ws2.append(["张三", "13800000000", 29.9])
    ws2.append(["李四", "13900000000", 199.0])

    wb.create_sheet("草稿")  # 空 Sheet

    wb.save(str(path))
    return str(path)
